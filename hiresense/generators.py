"""Templates and document generation (HTML preview, PDF, DOCX).

All three outputs come from one function, document_blocks(), so they always show the same content.
Changing templates changes presentation only.
"""

import html
import io
from xml.sax.saxutils import escape as xml_escape

from .resume_model import ENTRY_SECTIONS, PERSONAL_FIELDS, SECTION_ORDER, SECTION_TITLES

TEMPLATES = ("Classic", "Modern", "ATS-Friendly")
PURPLE = "#6d28d9"


def document_blocks(data: dict) -> dict:
    """Ordered, template-independent content.

    Returns {"name", "role", "contact", "header", "sections"}; sections is a list of (title, lines) where each
    line is (kind, text) or ("entry", heading, dates). Hidden sections are left out.
    """
    personal = data.get("personal", {})
    hidden = set(data.get("hidden_sections", []))
    header = [personal.get(field, "") for field in PERSONAL_FIELDS]
    contact = " | ".join(part for part in (personal.get("email"), personal.get("phone"), personal.get("location")) if part)
    sections = []
    for section in SECTION_ORDER:
        if section in hidden:
            continue
        lines = []
        if section == "summary":
            if data.get("summary"):
                lines.append(("para", data["summary"]))
        elif section in ENTRY_SECTIONS:
            for entry in data.get(section, []):
                head = " — ".join(part for part in (entry.get("title", ""), entry.get("organization", "")) if part)
                if head or entry.get("dates"):
                    lines.append(("entry", head, entry.get("dates", "")))
                details = [entry.get(key, "") for key in ("location", "field_of_study", "grade", "link")]
                details = [part for part in details if part]
                if details:
                    lines.append(("para", " · ".join(details)))
                for bullet in entry.get("bullets", []):
                    if bullet.get("text"):
                        lines.append(("bullet", bullet["text"]))
        elif section == "skills":
            groups: dict[str, list[str]] = {}
            for skill in data.get("skills", []):
                if skill.get("name"):
                    groups.setdefault(skill.get("category", ""), []).append(skill["name"])
            for category, names in groups.items():
                joined = ", ".join(names)
                lines.append(("para", f"{category}: {joined}" if category else joined))
        elif section == "links":
            for link in data.get("links", []):
                text = " ".join(part for part in (link.get("label", ""), link.get("url", "")) if part)
                if text:
                    lines.append(("bullet", text))
        else:
            for item in data.get(section, []):
                if item.get("text"):
                    lines.append(("bullet", item["text"]))
        if lines:
            sections.append((SECTION_TITLES[section], lines))
    return {"name": personal.get("name", ""), "role": personal.get("role", ""), "contact": contact,
            "header": header, "sections": sections}


# ---------- HTML preview ----------

_CSS = {
    "Classic": """
        .rs{font-family:Georgia,'Times New Roman',serif;color:#1f2937;background:#fff;padding:28px;border:1px solid #e5e7eb;}
        .rs .name{text-align:center;font-size:26px;font-weight:700;}
        .rs .contact{text-align:center;font-size:13px;color:#4b5563;margin-bottom:14px;}
        .rs h3{font-size:14px;text-transform:uppercase;letter-spacing:.05em;border-bottom:1px solid #1f2937;padding-bottom:3px;margin:16px 0 8px;}
        .rs .entry{display:flex;justify-content:space-between;gap:12px;font-weight:700;margin-top:8px;}
    .rs .entry .date{font-style:italic;font-weight:400;white-space:nowrap;}
    .rs .role{text-align:center;font-style:italic;font-size:14px;}
    .rs ul{margin:4px 0 0 18px;padding:0;} .rs p{margin:4px 0;font-size:14px;}
    """,
    "Modern": """
        .rs{font-family:Helvetica,Arial,sans-serif;color:#111827;background:#fff;padding:28px;border-left:6px solid #6d28d9;}
        .rs .name{font-size:28px;font-weight:800;color:#111827;}
        .rs .contact{font-size:13px;color:#6d28d9;margin-bottom:12px;}
        .rs h3{font-size:13px;color:#6d28d9;text-transform:uppercase;letter-spacing:.08em;margin:18px 0 6px;}
        .rs .entry{display:flex;justify-content:space-between;gap:12px;font-weight:600;margin-top:8px;}
        .rs .entry .date{font-weight:400;color:#6b7280;white-space:nowrap;}
        .rs ul{margin:4px 0 0 18px;padding:0;} .rs p{margin:4px 0;font-size:14px;}
    """,
    "ATS-Friendly": """
        .rs{font-family:Arial,sans-serif;color:#000;background:#fff;padding:28px;}
        .rs .name{font-size:20px;font-weight:700;} .rs .contact{font-size:13px;margin-bottom:10px;}
        .rs h3{font-size:14px;font-weight:700;margin:14px 0 4px;text-transform:uppercase;}
        .rs .entry{display:flex;justify-content:space-between;gap:12px;font-weight:700;margin-top:6px;}
        .rs .entry .date{font-weight:400;white-space:nowrap;}
        .rs ul{margin:4px 0 0 18px;padding:0;} .rs p{margin:3px 0;font-size:13px;}
    """,
}


def render_html(data: dict, template: str) -> str:
    """HTML preview. Every user value is escaped, so user input can never run as markup or script."""
    if template not in TEMPLATES:
        raise ValueError(f"Unknown template: {template}")
    blocks = document_blocks(data)
    parts = [f"<style>{_CSS[template]}</style>", '<div class="rs">']
    parts.append(f'<div class="name">{html.escape(blocks["name"] or "Your name")}</div>')
    if blocks["role"]:
        parts.append(f'<div class="role">{html.escape(blocks["role"])}</div>')
    if blocks["contact"]:
        parts.append(f'<div class="contact">{html.escape(blocks["contact"])}</div>')
    for title, lines in blocks["sections"]:
        parts.append(f"<h3>{html.escape(title)}</h3>")
        bullets_open = False
        for line in lines:
            kind, text = line[0], line[1]
            if kind == "bullet":
                if not bullets_open:
                    parts.append("<ul>")
                    bullets_open = True
                parts.append(f"<li>{html.escape(text)}</li>")
                continue
            if bullets_open:
                parts.append("</ul>")
                bullets_open = False
            if kind == "entry":
                date = html.escape(line[2]) if line[2] else ""
                parts.append(f'<div class="entry"><span>{html.escape(text)}</span><span class="date">{date}</span></div>')
            else:
                parts.append(f"<p>{html.escape(text)}</p>")
        if bullets_open:
            parts.append("</ul>")
    parts.append("</div>")
    return "".join(parts)


# ---------- PDF (ReportLab) ----------

def build_pdf(data: dict, template: str) -> bytes:
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib import colors
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, ListFlowable, ListItem

    if template not in TEMPLATES:
        raise ValueError(f"Unknown template: {template}")
    font = {"Classic": "Times-Roman", "Modern": "Helvetica", "ATS-Friendly": "Helvetica"}[template]
    bold = {"Classic": "Times-Bold", "Modern": "Helvetica-Bold", "ATS-Friendly": "Helvetica-Bold"}[template]
    accent = colors.HexColor(PURPLE) if template == "Modern" else colors.black
    base = getSampleStyleSheet()["Normal"]
    name_style = ParagraphStyle("name", parent=base, fontName=bold, fontSize=20, leading=24,
                                alignment=TA_CENTER if template == "Classic" else 0, textColor=colors.black)
    contact_style = ParagraphStyle("contact", parent=base, fontName=font, fontSize=9, leading=12,
                                   alignment=TA_CENTER if template == "Classic" else 0, textColor=colors.HexColor("#4b5563"))
    heading_style = ParagraphStyle("heading", parent=base, fontName=bold, fontSize=11, leading=14,
                                   spaceBefore=10, spaceAfter=3, textColor=accent)
    entry_style = ParagraphStyle("entry", parent=base, fontName=bold, fontSize=10, leading=13, spaceBefore=4)
    body_style = ParagraphStyle("body", parent=base, fontName=font, fontSize=10, leading=13)

    blocks = document_blocks(data)
    story = [Paragraph(xml_escape(blocks["name"] or "Your name"), name_style)]
    if blocks["role"]:
        story.append(Paragraph(xml_escape(blocks["role"]), contact_style))
    if blocks["contact"]:
        story.append(Paragraph(xml_escape(blocks["contact"]), contact_style))
    for title, lines in blocks["sections"]:
        story.append(Paragraph(xml_escape(title.upper() if template != "Modern" else title), heading_style))
        pending_bullets: list = []

        def flush():
            if pending_bullets:
                story.append(ListFlowable(
                    [ListItem(Paragraph(xml_escape(text), body_style), leftIndent=12) for text in pending_bullets],
                    bulletType="bullet", start="•", leftIndent=12,
                ))
                pending_bullets.clear()

        for line in lines:
            kind = line[0]
            if kind == "bullet":
                pending_bullets.append(line[1])
                continue
            flush()
            if kind == "entry":
                label = " — ".join(part for part in (line[1], line[2]) if part)
                story.append(Paragraph(xml_escape(label), entry_style))
            else:
                story.append(Paragraph(xml_escape(line[1]), body_style))
        flush()
    story.append(Spacer(1, 4))

    buffer = io.BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=48, rightMargin=48, topMargin=44, bottomMargin=44,
                                 title=blocks["name"] or "Resume", author="HireSense AI")
    document.build(story)
    return buffer.getvalue()


# ---------- DOCX (python-docx) ----------

def build_docx(data: dict, template: str) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    if template not in TEMPLATES:
        raise ValueError(f"Unknown template: {template}")
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman" if template == "Classic" else "Arial"
    style.font.size = Pt(10.5)
    blocks = document_blocks(data)

    name_para = document.add_paragraph()
    name_run = name_para.add_run(blocks["name"] or "Your name")
    name_run.bold = True
    name_run.font.size = Pt(18)
    if template == "Classic":
        name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if blocks["role"]:
        role_para = document.add_paragraph(blocks["role"])
        role_para.alignment = WD_ALIGN_PARAGRAPH.CENTER if template == "Classic" else WD_ALIGN_PARAGRAPH.LEFT
    if blocks["contact"]:
        contact_para = document.add_paragraph(blocks["contact"])
        if template == "Classic":
            contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for title, lines in blocks["sections"]:
        heading = document.add_paragraph()
        run = heading.add_run(title.upper() if template != "Modern" else title)
        run.bold = True
        run.font.size = Pt(11.5)
        if template == "Modern":
            run.font.color.rgb = RGBColor(0x6D, 0x28, 0xD9)
        for line in lines:
            kind = line[0]
            if kind == "bullet":
                document.add_paragraph(line[1], style="List Bullet")
            elif kind == "entry":
                entry_para = document.add_paragraph()
                entry_run = entry_para.add_run(" — ".join(part for part in (line[1], line[2]) if part))
                entry_run.bold = True
            else:
                document.add_paragraph(line[1])

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
